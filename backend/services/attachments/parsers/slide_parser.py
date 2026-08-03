"""
Slide Parser — extracts text from PPTX files.
Handles slide title, body text, and speaker notes.
"""
from __future__ import annotations
from pathlib import Path
from typing import Tuple


def parse_slides(file_path: str) -> Tuple[str, int]:
    """
    Returns (full_text, slide_count).
    """
    try:
        from pptx import Presentation
        from pptx.util import Pt

        prs    = Presentation(file_path)
        slides = []
        count  = len(prs.slides)

        for i, slide in enumerate(prs.slides):
            parts = [f"--- Slide {i + 1} ---"]

            # Slide title
            if slide.shapes.title and slide.shapes.title.text.strip():
                parts.append(f"Title: {slide.shapes.title.text.strip()}")

            # All text frames
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                if shape == slide.shapes.title:
                    continue  # already handled
                texts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if texts:
                    parts.append("\n".join(texts))

            # Speaker notes
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    parts.append(f"[Notes: {notes_text}]")

            slides.append("\n".join(parts))

        return "\n\n".join(slides), count

    except ImportError:
        return "[PPTX Error: python-pptx not installed]", 0
    except Exception as e:
        return f"[PPTX Parse Error: {e}]", 0
