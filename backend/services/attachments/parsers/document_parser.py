"""
Document Parser — extracts text from DOCX, TXT, MD, RTF files.
"""
from __future__ import annotations
from pathlib import Path
from typing import Tuple


def parse_document(file_path: str, category: str = "document") -> Tuple[str, int]:
    """
    Returns (full_text, page_count_estimate).
    page_count is estimated as text_length / 3000 chars per page.
    """
    ext = Path(file_path).suffix.lower()

    if ext in (".docx", ".doc"):
        return _parse_docx(file_path)
    elif ext == ".rtf":
        return _parse_rtf(file_path)
    elif ext in (".txt", ".md", ".rst", ".log"):
        return _parse_plaintext(file_path)
    else:
        # Generic: try UTF-8 read
        return _parse_plaintext(file_path)


def _parse_docx(file_path: str) -> Tuple[str, int]:
    try:
        from docx import Document
        doc   = Document(file_path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract tables
        table_texts = []
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append(" | ".join(c.text.strip() for c in row.cells))
            if rows:
                table_texts.append("\n".join(rows))

        full = "\n\n".join(paras)
        if table_texts:
            full += "\n\n--- Tables ---\n" + "\n\n".join(table_texts)

        # Extract properties
        try:
            props = doc.core_properties
            header = ""
            if props.title:
                header = f"Title: {props.title}\n"
            if props.author:
                header += f"Author: {props.author}\n"
            if header:
                full = header + "\n" + full
        except Exception:
            pass

        pages = max(1, len(full) // 3000)
        return full, pages
    except ImportError:
        return "[DOCX Error: python-docx not installed]", 0
    except Exception as e:
        return f"[DOCX Parse Error: {e}]", 0


def _parse_rtf(file_path: str) -> Tuple[str, int]:
    try:
        from striprtf.striprtf import rtf_to_text
        raw  = Path(file_path).read_text(encoding="utf-8", errors="replace")
        text = rtf_to_text(raw).strip()
        return text, max(1, len(text) // 3000)
    except ImportError:
        # Fallback: strip RTF tags with regex
        import re
        raw  = Path(file_path).read_text(encoding="utf-8", errors="replace")
        text = re.sub(r"\{.*?\}|\\[a-z]+\d* ?|[{}\\]", " ", raw).strip()
        return text, max(1, len(text) // 3000)
    except Exception as e:
        return f"[RTF Parse Error: {e}]", 0


def _parse_plaintext(file_path: str) -> Tuple[str, int]:
    try:
        import chardet
        raw      = Path(file_path).read_bytes()
        detected = chardet.detect(raw)
        enc      = detected.get("encoding") or "utf-8"
        text     = raw.decode(enc, errors="replace").strip()
        return text, max(1, len(text) // 3000)
    except Exception:
        try:
            text = Path(file_path).read_text(encoding="utf-8", errors="replace").strip()
            return text, max(1, len(text) // 3000)
        except Exception as e:
            return f"[Text Parse Error: {e}]", 0
